import random
import openpyxl
import pyttsx3
import speech_recognition as sr
import colorama
from fuzzywuzzy import fuzz                         #Расстояние Левенштейна для нечеткого сравнения строк
import datetime
from os import system, path, makedirs, startfile
import sys
from random import choice
from pyowm import OWM                               #Open Weather Map
from pyowm.utils.config import get_default_config
import webbrowser
import configparser
from psutil import virtual_memory as memory
import nltk
import wikipedia as wiki
import requests
from bs4 import BeautifulSoup
import psutil
import winsound
import datefinder
from pywhatkit import playonyt
from PyQt5 import QtWidgets, QtCore
import interface
import threading
import re
import os
from os.path import join
import io
from docx import Document
from vosk import Model, KaldiRecognizer
import pyaudio
import json
from googletrans import Translator

with open('BASE_INTENTS.json', 'r') as jsn:
    BASE_INTENTS = json.load(jsn)


class Assistant(QtWidgets.QMainWindow, interface.Ui_MainWindow, threading.Thread):
    settings = configparser.ConfigParser()
    settings.read('settings.ini')

    config_dict = get_default_config()  # Инициализация get_default_config()
    config_dict['language'] = 'ru'  # Установка языка

    last_dir = ''

    def __init__(self):
        super().__init__()
        self.setupUi(self)

        self.model = Model(r"C:\Users\mrgod\PycharmProjects\VoiceAsistent\vosk-model-small-ru-0.22")  # полный путь к модели
        self.rec = KaldiRecognizer(self.model, 16000)
        self.p = pyaudio.PyAudio()
        self.stream = None
        self.text = ""

        self.pushButton.clicked.connect(self.start_thread)
        self.pushButton_2.clicked.connect(self.stop)
        self.working = False
        self.r = sr.Recognizer()
        self.text = ''

        self.cmds = {
            ('текущее время', 'сейчас времени', 'который час'): self.time,
            ('привет', 'добрый день', 'здравствуй'): self.hello,
            ('пока', 'вырубись'): self.quite,
            ('выключи компьютер', 'выруби компьютер'): self.shut,
            ('какая погода', 'погода', 'погода на улице'): self.weather,
            ('добавить задачу', 'добавить заметку', 'создай заметку', 'создай задачу'): self.task_planner,
            ('список задач', 'список заметок', 'задачи', 'заметки'): self.task_list,
            ('удалить задачу', 'удали задачу', 'хочу удалить задачу'): self.task_delete,
            ('загруженость компьютера', 'загруженость системы', 'загруженость','состояние системы', 'какая загрузка системы', 'какая загрузка'): self.check_memory,
            ('включи музыку', 'вруби музон', 'вруби музыку', 'включи музон', 'врубай музыку'): self.music,
            ('расскажи анекдот', 'анекдот', 'пошути'): self.joke,
            ('создай папку', 'создать папку'): self.create_folder,
            ('создай файл', 'создать файл'): self.create_file,
            ('какой курс валют', 'скажи курс валют', 'курс валют'): self.currency,
            ('место на диске', 'сколько памяти', 'сколько памяти на диске', 'сколько места'): self.disk_usage,
            ('перезагрузи компьютер', 'перезагрузи комп', 'перезагружай комп', 'перезагрузи'): self.restart_pc,
            ('забавный факт', 'смешной факт', 'факт дня', 'интересный факт'): self.facts,
            ('сыграем в камень ножницы бумага', 'давай сыграем в камень ножницы бумага', 'камень ножницы бумага','сыграем в цу е фа','сыграем в цу-е-фа'): self.tsu_e_fa,
            ('сыграем в угадай число', 'давай сыграем в угадай число', 'угадай число'): self.guess_number,
        }

        self.ndels = ['давид', 'дэвид', 'не мог бы ты']

        self.commands = [
            'текущее время', 'сейчас времени', 'который час',
            'открой браузер', 'открой интернет', 'запусти браузер',
            'привет', 'добрый день', 'здравствуй',
            'пока', 'вырубись',
            'выключи компьютер', 'выруби компьютер',
            'какая погода', 'погода', 'погода на улице', 'какая погода на улице',
            'добавить задачу', 'добавить заметку', 'создай заметку', 'создай задачу','список задач',
            'список заметок', 'задачи','заметки',
            'загруженость компьютера', 'загруженость системы', 'какая загрузка',
            'включи музыку', 'вруби музон', 'вруби музыку', 'включи музон', 'врубай музыку',
            'расскажи анекдот', 'анекдот', 'пошути'
            'место на диске', 'сколько памяти', 'сколько памяти на диске', 'сколько места',
            'перезагрузи компьютер', 'перезагрузи комп', 'перезагружай комп', 'перезагрузи',
            'забавный факт', 'смешной факт', 'факт дня', 'интересный факт','удалить задачу', 'удали задачу',
            'хочу удалить задачу' 'создай файл', 'создать файл', 'сыграем в камень ножницы бумага', 'давай сыграем в камень ножницы бумага',
            'камень ножницы бумага', 'сыграем в угадай число', 'давай сыграем в угадай число', 'угадай число'
        ]

        self.num_task = 0
        self.j = 0
        self.ans = ''

        wiki.set_lang('ru')

    def start_stream(self):
        self.stream = self.p.open(
            format=pyaudio.paInt16,
            channels=1,
            rate=16000,
            input=True,
            frames_per_buffer=16000
        )
        self.stream.start_stream()

    def stop_stream(self):
        self.stream.stop_stream()
        self.stream.close()

    def cleaner(self, text):
        self.text = text

        for i in self.ndels:
            self.text = self.text.replace(i, '').strip()
            self.text = self.text.replace('  ', ' ').strip()

        self.ans = self.text

        for i in range(len(self.commands)):
            k = fuzz.ratio(text, self.commands[i])
            if (k > 70) & (k > self.j):
                self.ans = self.commands[i]
                self.j = k

        return str(self.ans)

    def check_memory(self):
        mem = memory()
        self.talk(f'Компьютер загружен на {round(mem.percent)}%')

    def intent_cleaner(self, text): #Функция очистки текста
        cleaned_text = ''
        for i in text.lower(): #Текст в нижнем регистре
            if i in 'абвгдеёжзийклмнопрстуфхцчшщъыьэюяabcdefghijklmnopqrstuvwxyz ':
                cleaned_text += i
        return cleaned_text

    #Расстояние Левенштейна используется для нечеткого сравнения строк

    def match(self, text, example):  #Функция для сравнения текста с тем что есть в базе интентов
        return nltk.edit_distance(text, example) / len(example) < 0.4 if len(example) > 0 else False
    # С помощью растояния Левенштейна проверяем на сколько написанный пользователем текст совпадает с намерением в базе интентов и выдаем ответ

    def get_intent(self, text): #Определяет интент в тексте
        for intent in BASE_INTENTS['intents']:
            if 'examples' in BASE_INTENTS['intents'][intent]:
                for example in BASE_INTENTS['intents'][intent]['examples']:
                    if self.match(self.intent_cleaner(text), self.intent_cleaner(example)):
                        return intent

    def intenter(self, text):
        intent = self.get_intent(text) #Пытаемся понять намерения и сравнить по Левенштейну

        if intent is None:
            return self.talk(" ") #Тест если нету такого интента

        self.talk(choice(BASE_INTENTS['intents'][intent]['responses']))

    def task_planner(self):
        self.talk("Что добавить в список задач?")
        task = self.listen()

        with open(r'C:\Users\mrgod\PycharmProjects\VoiceAsistent\TODO_LIST.txt', 'a') as file:
            file.write(f'{task}\n')

        self.talk(f'Задача {task} добавлена в список задач!')

    def task_list(self):
        with open(r'C:\Users\mrgod\PycharmProjects\VoiceAsistent\TODO_LIST.txt', 'r') as file:
            tasks = file.read()
        self.talk(f"Список задач:\n{tasks}")

    def task_delete(self):
        self.talk("Какую задачу удалить?")
        task = self.listen()

        with open(r'C:\Users\mrgod\PycharmProjects\VoiceAsistent\TODO_LIST.txt') as file:
            lines = file.readlines()

        pattern = re.compile(re.escape(task))                                                   #Проходим по массиву строк и сравниваем с искомой строкой
        with open(r'C:\Users\mrgod\PycharmProjects\VoiceAsistent\TODO_LIST.txt', 'w') as file:  #Если строка не найдена записываем строку в файл
            for line in lines:
                result = pattern.search(line)
                if result is None:
                    file.write(line)
                    self.talk(f'Задача {task} удалена')

    def wiki_search(self, task):
        try:
            info = wiki.summary(task, sentences=3)
            self.talk((info.replace('англ', '')).replace('род.', 'родился').replace('(.', '').replace(')', '')
                      .replace(';', '').replace("(урожд. —", "").replace("урожд.", "").replace("  ", " "))
        except wiki.exceptions.PageError:
            pass
        except wiki.exceptions.WikipediaException:
            pass

    def music(self):
        self.talk(choice(['Приятного прослушивания!']))
        music_list = ['https://www.youtube.com/watch?v=HKcsaovoB6s', 'https://vk.com/audios205597577?section=general',
                      'https://www.youtube.com/watch?v=GQTFPIoz7Uc&t=4106s', 'https://www.youtube.com/watch?v=8wy4ACtLlfQ&t=2235s']
        webbrowser.open(choice(music_list))

    def translate(self, task):
        self.talk(choice(['Сейчас попытаюсь перевести', 'Сейчас переведу', 'Сейчас попробую перевести']))
        variants = ['переведи', 'перевести', 'перевод', 'перевести', 'переводить']
        for i in variants:
            print(i)
            task = task.replace(i, '').replace('  ', ' ')
        print(task)
        webbrowser.open(f'https://translate.yandex.by/?source_lang=ru&target_lang=en&text={task}')

    def search_on_map(self, task):
        for j in ('найди на карте', 'найти на карте', 'поищи на карте'):
            task = task.replace(j, '').replace('  ', ' ').strip()
            for i in ['находится', 'расположен', 'где']:
                for k in ['мне', 'нам', 'всем', 'им']:
                    if fuzz.ratio(task.split()[0], i) > 70 or fuzz.ratio(task.split()[0], k) > 70:
                        task = ' '.join(task.split()[1:])
                    task = task.replace(i, '').replace(k, '').replace('  ', ' ').strip()
        webbrowser.open(f'https://google.com/maps/search/{task}')
        self.talk(choice(['Сейчас всё найду', 'Ищу на карте', 'Уже ищу']))

    def web_search(self, search):
        mode = ''

        youtube = [
            'найди видео', 'найди на ютубе', 'найти на ютубе', 'найди на youtube',
            'найти видео', 'ищи видео', 'ищи на ютубе', 'ютуб', 'ютюбе', 'найди на ютюбе' 'видео']

        words = ['найди', 'найти', 'ищи', 'кто такой', 'что такое']
        remove = ["пожалуйста", "ладно", "давай", "сейчас"]

        if mode == '':
            for i in range(len(youtube)):
                if youtube[i] in search:
                    mode = 'youtube'

        if mode == '':
            for i in range(len(words)):
                if words[i] in search:
                    mode = 'web'

        if mode == 'youtube':
            for i in youtube:
                search = search.replace(i, '')
                for j in remove:
                    search = search.replace(j, '')
                    search = search.strip()

            self.talk(f'Ищу {search} на ютубе')
            webbrowser.open(f'https://www.youtube.com/results?search_query={search}')

            self.talk("Желаете сразу открыть первое видео?")
            self.text = self.listen()
            print(self.text)
            if (fuzz.ratio(self.text, 'давай') > 60) or (fuzz.ratio(self.text, 'да') > 60) \
                    or (fuzz.ratio(self.text, "открой") > 60) or (fuzz.ratio(self.text, "открывай") > 60):
                self.talk('Открываю видео')
                playonyt(search)
            elif fuzz.ratio(self.text, 'нет') > 60 or (fuzz.ratio(self.text, "не надо") > 60):
                self.talk("Как пожелаете")
        elif mode == 'web':
            for i in words:
                search = search.replace(i, '')
                for j in remove:
                    search = search.replace(j, '')
                    search = search.strip()
            self.talk(f"Ищу {search}")
            webbrowser.open(f'https://www.google.com/search?q={search}&oq={search}'
                            f'81&aqs=chrome..69i57j46i131i433j0l5.2567j0j7&sourceid=chrome&ie=UTF-8')

            self.wiki_search(search)

    def recognizer(self):
        self.text = self.cleaner(self.listen())
        print(self.text)

        if self.text.startswith(('открой', 'запусти', 'зайди', 'зайди на')):
            self.opener(self.text)

        elif self.text.startswith(('найди на карте', 'найти на карте', 'поищи на карте')):
            self.search_on_map(self.text)

        elif self.text.startswith(('найди', 'найти', 'ищи', 'кто такой', 'что такое')):
            self.web_search(self.text)

        elif self.text.startswith(('переведи', 'перевести', 'перевод', 'перевести', 'переводить')):
            self.translate(self.text)

        elif self.text.startswith(('поставь будильник', 'поставить будильник', 'запусти будильник', 'запустить будильник', 'будильник')):
            self.alarmclock(self.text)

        for tasks in self.cmds:                            #Сверяет сказанное с командами которые в cmds
            for task in tasks:
                if fuzz.ratio(task, self.text) >= 80:
                    self.cmds[tasks]()
                    return

        self.intenter(self.text)

    def facts(self):
        fact_url = "https://randstuff.ru/fact/"
        response = requests.get(fact_url)
        soup = BeautifulSoup(response.content, 'html.parser').findAll('td')
        items = list(soup)
        funny_fact = items[0]
        self.talk(str(funny_fact).replace('<td>', '').replace('</td>', '').replace('  ', ' '))

    def disk_usage(self):
        total, used, free, percent = psutil.disk_usage('C:/')
        self.talk(f"Всего {total // (2 ** 30)} гигабайт, используется {used // (2 ** 30)} гигабайт,"
                  f"свободно {free // (2 ** 30)}, Системный диск используется на {percent} процентов")

    def create_folder(self):
        # Получаем путь к рабочему столу
        desktop_path = (r"C:\Users\mrgod\OneDrive\Рабочий стол")
        self.talk("Как назовем папку?")
        name = self.listen()
        # Создаем полный путь к новой папке на рабочем столе
        new_folder_path = os.path.join(desktop_path, name)

        # Проверяем, существует ли папка уже
        if os.path.exists(new_folder_path):
            self.talk("Папка уже существует")
        else:
            # Создаем папку на рабочем столе
            os.mkdir(new_folder_path)
            self.talk("Папка успешно создана")

    def create_file(self):
        desktop_path = (r"C:\Users\mrgod\OneDrive\Рабочий стол")
        self.talk("В какой папке создать файл")
        folder = self.listen()

        # Создаем полный путь к новой папке на рабочем столе
        new_folder_path = os.path.join(desktop_path, folder)

        # Проверяем, существует ли папка уже
        if os.path.exists(new_folder_path):
            while True:
                self.talk("Какого типа файл вы хотите создать?\nТекстовый\nЭксель\nВорд")
                name = self.listen()
                if name == "текстовый" or name == "текст":
                    self.create_text_file(new_folder_path)
                    break
                elif name == "эксель" or name == "эксэль" or name == "excel":
                    self.create_exel_file(new_folder_path)
                    break
                elif name == "ворд" or name == "ворт" or name == "word" or name == "world":
                    self.create_word_file(new_folder_path)
                    break

        else:
            self.talk("Такой папки нету,\n Вначале создайте ее\n с помощью голосовой команды")

    def create_text_file(self, file_path):
        while True:
            self.talk("как назовем")
            name = self.listen()
            new_file_path = os.path.join(file_path, name)
            new_path = f"{new_file_path}.txt"

            if os.path.exists(new_path):
                self.talk("Файл уже существует, придумайте новое название")
            else:
                with open(new_path, "w") as file:
                    self.talk("хотите туда что-то записать?")
                    choice = self.listen()
                    if choice == "да":
                        self.talk("что именно")
                        text = self.listen()
                        file.write(text)
                    elif choice == "нет":
                        self.talk("Ну как хотите")
                self.talk(f"Файл {name} успешно создан")
                break

    def create_exel_file(self, file_path):
        self.talk("как назовем")
        name = self.listen()
        new_file_path = os.path.join(file_path, name)
        new_path = f"{new_file_path}.xlsx"
        self.talk(new_path)
        wb = openpyxl.Workbook()
        wb.save(new_path)
        self.talk(f"Файл Excel успешно создан.")

    def create_word_file(self, file_path):
        self.talk("как назовем")
        name = self.listen()
        new_file_path = os.path.join(file_path, name)
        new_path = f"{new_file_path}.docx"
        if os.path.exists(new_path):
            self.talk("Файл уже существует.")
        else:
            document = Document()
            document.save(new_path)
            self.talk("хотите туда что-то записать?")
            choice = self.listen()
            if choice == "да":
                self.talk("что именно")
                text = self.listen()
                document.add_paragraph(text)
                document.save(new_path)
            elif choice == "нет":
                self.talk("Ну как хотите")
            self.talk(f"Файл {name} успешно создан")

    def currency(self):
        c = Currency()
        choice_usd = ['к доллару', 'доллару', 'давай доллару']
        choice_rub = ['к рублю', 'рублю', 'к российскому рублю', 'российскому рублю']
        choice_exit = ['не хочу', 'больше не хочу узнавать', 'хватит', 'не надо', 'я передумал']
        self.talk("К какой валюте хотите узнать курс? (к доллару или к российскому рублю?)")
        choice = self.listen()
        for i in choice_exit:
            if i == choice:
                self.talk("Ну как хотите")
        for i in choice_rub:
            if i == choice:
                self.talk(fr"1 российский рубль  {c.currency_rub()} белларусских рублей")
        for i in choice_usd:
            if i == choice:
                self.talk(fr"1 доллар - {c.currency_usd()} беларусских рублей")

    def joke(self):
        link = requests.get('http://anekdotme.ru/random')
        parse = BeautifulSoup(link.text, "html.parser")
        select = parse.select('.anekdot_text')
        get = (select[0].getText().strip())
        reg = re.compile('[^a-zA-Zа-яА-я ^0-1-2-3-4-5-6-7-8-9:.,!?-]')
        joke = reg.sub('', get)
        self.talk(joke)

    def time(self):
        now = datetime.datetime.now()
        self.talk(f"Сейчас {now.hour} : {now.minute}")

    def opener(self, task):
        links = {
            ('youtube', 'ютуб', 'ютюб'): 'https://youtube.com/',
            ('вк', 'вконтакте', 'контакт', 'vk'): 'https:vk.com/',
            ('insta', 'instagram', 'инста', 'инсту'): 'https://www.instagram.com/',
            ('гугл почту', 'gmail', 'гмейл', 'гмеил', 'гмаил'): 'http://gmail.com/',
            ('яндекс почту', 'почту'): 'https://mail.yandex.by/',
            ('яндекс', 'браузер', 'интернет', 'browser'): 'https://yandex.by/',
            ('твич', 'твитч', 'twitch'): 'https://www.twitch.tv/',
        }

        def open_discord():
            sp = Search_program()
            name = 'Discord.exe'
            path_to_program = sp.search_in_txt(name)
            try:
                startfile(path_to_program)
                system("cls")
                self.talk(choice(["Открываю дискорд", "Включаю дискорд", "запускаю дискорд"]))
            except:
                self.talk("Я искал эту программу на компьютере, скажите еще раз мне открыть программу ")

        def open_tg():
            sp = Search_program()
            name = 'Telegram.exe'
            path_to_program = sp.search_in_txt(name)
            try:
                startfile(path_to_program)
                self.talk(choice(["Открываю телеграм", "Включаю телеграм", "запускаю телеграм"]))
            except:
                self.talk("Я искал эту программу на компьютере, скажите еще раз мне открыть программу ")

        def open_steam():
            sp = Search_program()
            name = 'steam'
            path_to_program = sp.search_in_txt(name)
            try:
                startfile(path_to_program)
                self.talk(choice(["Открываю стим", "Открываю стим, но много не играйте", "Запуская стим, хорошей игры"]))
            except:
                self.talk("Я искал эту программу на компьютере, скажите еще раз мне открыть программу ")


        programs = {
            'discord': open_discord, 'дискорд': open_discord, 'дискорт': open_discord, 'дис корд': open_discord, 'дис корт': open_discord, 'дис скорд': open_discord, 'дис скорт': open_discord,
            'телеграм': open_tg, 'telegram': open_tg,
            'стим': open_steam, 'steam': open_steam, 'стин': open_steam, 'стеам': open_steam,
        }

        j = 0
        if 'и' in task:
            task = task.replace('и', '').replace('  ', ' ')
        double_task = task.split()
        if j != len(double_task):
            for i in range(len(double_task)):
                for vals in links:
                    for word in vals:
                        if fuzz.ratio(word, double_task[i]) > 75:
                            webbrowser.open(links[vals])
                            self.talk('Открываю ' + double_task[i])
                            j += 1
                            break
                else:
                    for vals in programs:
                        if fuzz.ratio(vals, double_task[i]) > 75:
                            programs[vals]()
                            j += 1
                            break

    def tsu_e_fa(self):
        self.talk("Ну давайте сыграем")

        while True:
            possible_actions = ["камень", "бумага", "ножницы"]
            computer_action = random.choice(possible_actions)
            self.talk("Я выбрал, выбирайте вы")
            user_action = self.listen()

            if user_action == "не хочу" or user_action == "я передумал":
                self.talk("Ну как хотите")
            if user_action == computer_action:
                self.talk(f"Мы оба выбрали {user_action}, ничья")
            elif (
                (user_action == "камень" and computer_action == "ножницы") or
                (user_action == "бумага" and computer_action == "камень") or
                (user_action == "ножницы" and computer_action == "бумага")
            ):
                self.talk("Вы выиграли!")
            else:
                self.talk("Вы проиграли!")
            self.talk("Сыграем еще?")
            play_again = self.listen()
            if play_again.lower() != "да":
                self.talk("Ну как хотите")
                break

    def guess_number(self):
        NumberToGuess = str(random.randint(1, 10))
        self.talk("Ну давайте сыграем, только давайте до 10-ти, а то мы так долго будем")
        self.talk("кто загадывает?")
        num = []
        who_play = self.listen()

        if who_play == "хватит" or who_play == "не хочу":
            self.talk("Ну как хотите")

        if who_play == "я":
            self.talk("Ну загадывай")
            while True:
                userGuess = self.listen()
                if int(userGuess) > 10 or int(userGuess) < 1:
                    self.talk("Я же сказал от 1 до 10")
                else:
                    while True:
                        NumberToGuessBot = str(random.randint(1, 10))
                        while num.__contains__(NumberToGuessBot):
                            NumberToGuessBot = str(random.randint(1, 10))
                        num.append(NumberToGuessBot)

                        self.talk(f"Я думаю {NumberToGuessBot}")
                        answer = self.listen()
                        if userGuess == "хватит" or userGuess == "не хочу":
                            self.talk("Ну как хотите")
                            break
                        if answer == "да":
                            if userGuess != NumberToGuessBot:
                                self.talk(f"Дак ты не это загадывал, я же все знаю, твое число было {userGuess}")
                                break
                            self.talk("Было легко")
                            break
                        if answer == "нет":
                            if NumberToGuessBot == userGuess:
                                self.talk("Не ври, ты это число загадал")
                                break
                    break

        if who_play == "ты" or who_play == "вы":
            self.talk("Я загадал, угадывай")
            userChoice = self.listen()

            while userChoice != NumberToGuess:
                if userChoice == "хватит" or userChoice == "не хочу":
                    self.talk("Ну как хотите")
                    break
                if userChoice > NumberToGuess:
                    self.talk("Число должно быть меньше!")
                elif userChoice < NumberToGuess:
                    print("Число должно быть больше!")
                else:
                    print(f"Вы угадали, это число {NumberToGuess}")
                break

    def quite(self):
        self.talk(choice(['Надеюсь я смог вам чем-то помочь', 'Рад был помочь', 'Пока', 'Я пошел']))
        self.engine.stop()
        self.working = False
        system('cls')
        sys.exit(0)

    def restart_pc(self):
        self.talk("Подтвердите действие!")
        text = self.listen()
        print(text)
        if (fuzz.ratio(text, 'подтвердить') > 60) or (fuzz.ratio(text, "подтверждаю") > 60):
            self.talk('Действие подтверждено')
            self.talk('До скорых встреч!')
            system('shutdown /r /f /t 10 /c "Перезагрузка будет выполнена через 10 секунд"')
            self.quite()
        elif fuzz.ratio(text, 'отмена') > 60:
            self.talk("Действие не подтверждено")
        else:
            self.talk("Действие не подтверждено")

    def shut(self):
        self.talk("Подтвердите действие!")
        text = self.listen()
        print(text)
        if (fuzz.ratio(text, 'подтвердить') > 60) or (fuzz.ratio(text, "подтверждаю") > 60):
            self.talk('Действие подтверждено')
            self.talk('До скорых встреч!')
            system('shutdown /s /f /t 10')
            self.quite()
        elif fuzz.ratio(text, 'отмена') > 60:
            self.talk("Действие не подтверждено")
        else:
            self.talk("Действие не подтверждено")

    def hello(self):
        self.talk(choice(['Привет, чем могу помочь?', 'Здраствуйте', 'Приветствую']))

    def weather(self):
        self.talk("В каком городе хотите узнать погоду?")
        place = self.listen()

        self.talk("Какая это страна??")
        country = self.listen()

        try:
            country_and_place = place + ", " + country  # Запись города и страны в одну переменную через запятую
            owm = OWM('12ba88a83351c9466ed2a48a06bdd2c8')  # Ваш ключ с сайта open weather map
            mgr = owm.weather_manager()  # Инициализация owm.weather_manager()
            observation = mgr.weather_at_place(country_and_place)

            # Инициализация mgr.weather_at_place() И передача в качестве параметра туда страну и город

            w = observation.weather

            status = w.detailed_status  # Узнаём статус погоды в городе и записываем в переменную status
            w.wind()  # Узнаем скорость ветра
            humidity = w.humidity  # Узнаём Влажность и записываем её в переменную humidity
            temp = w.temperature('celsius')[
                'temp']  # Узнаём температуру в градусах по цельсию и записываем в переменную temp
            self.talk("В городе " + str(place) + " сейчас " + str(status) +  # Выводим город и статус погоды в нём
                        "\nТемпература " + str(
                round(temp)) + " градусов по цельсию" +  # Выводим температуру с округлением в ближайшую сторону
                        "\nВлажность составляет " + str(humidity) + "%" +  # Выводим влажность в виде строки
                        "\nСкорость ветра " + str(w.wind()['speed']) + " метров в секунду")   # Узнаём и выводим скорость ветра
        except:
            self.talk("Такого города или страны нету")

    def talk(self, text):
        self.engine = pyttsx3.init(debug=True)
        print(text)
        item = QtWidgets.QListWidgetItem()
        item.setTextAlignment(QtCore.Qt.AlignLeft)
        item.setText('DAVID:' + '\n' + text)

        self.listWidget.addItem(item)
        self.listWidget.scrollToBottom()
        self.engine.say(text)
        self.engine.runAndWait()
        self.engine.stop()

    def listen(self):
        self.start_stream()
        while True:
            print(f"{colorama.Fore.LIGHTGREEN_EX}Я вас слушаю...")
            data = self.stream.read(8000)
            if len(data) == 0:
                break
            if self.rec.AcceptWaveform(data):
                result = json.loads(self.rec.Result())
                self.text = result['text'].lower()

                if self.text != '':
                    item = QtWidgets.QListWidgetItem()
                    item.setTextAlignment(QtCore.Qt.AlignRight)
                    item.setText('Вы:' + '\n' + self.text)
                    self.listWidget.addItem(item)
                    self.listWidget.scrollToBottom()

                self.stop_stream()
                return self.text

    def alarmclock(self, text):
        time = datefinder.find_dates(text)
        date_and_time = ''
        for date_and_time in time:
            print(date_and_time)
        date_time = str(date_and_time)
        only_time = date_time[11:]
        print(only_time)
        self.talk(f'Будильник установлен на {only_time}')
        hour_time = int(only_time[:-6])
        min_time = int(only_time[3:-3])

        count = 0

        while True:
            if hour_time == datetime.datetime.now().hour:
                if min_time <= datetime.datetime.now().minute:
                    if count < 1:
                        print('Будильник сработал!')
                        winsound.PlaySound('sound.wav', winsound.SND_ALIAS)
                        count += 1
                    else:
                        break
                else:
                    self.recognizer()
            else:
                self.recognizer()

    def start_thread(self):
        self.hello()
        self.working = True
        self.thread = threading.Thread(target=self.main)
        self.thread.start()

    def stop(self):
        self.working = False
        self.quite()

    def main(self):

        while self.working:
            try:
                self.recognizer()
            except Exception as ex:
                print(ex)

class Currency():
    @classmethod
    def currency_soup(cls, currency_url):
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/110.0.0.0 Safari/537.36'}
        full_page = requests.get(currency_url, headers=headers)
        return BeautifulSoup(full_page.content, 'html.parser')

    @classmethod
    def currency_usd(cls):
        DOLLAR_BUN = 'https://www.google.by/search?q=курс+доллара+&sxsrf=APwXEdfluDW-LRIG3J2YfnMV_D7XsxuX9A%3A1684086414892&source=hp&ei=jh5hZM-eM8aM9u8P9eyzyAk&iflsig=AOEireoAAAAAZGEsnhZVb1-dJhVLyB300sQf9lSQkzzM&ved=0ahUKEwjPnZzkrvX-AhVGhv0HHXX2DJkQ4dUDCAk&uact=5&oq=курс+доллара+&gs_lcp=Cgdnd3Mtd2l6EAMyCAgAEIAEELEDMggIABCABBCxAzILCAAQgAQQsQMQgwEyCwgAEIAEELEDEIMBMgsIABCABBCxAxCDATILCAAQgAQQsQMQgwEyCAgAEIAEELEDMgUIABCABDILCAAQgAQQsQMQgwEyBQgAEIAEOgcIIxDqAhAnOggIABCPARDqAjoICC4QjwEQ6gI6CwgAEIoFELEDEIMBOg4ILhCKBRCxAxCDARDUAjoOCAAQgAQQsQMQgwEQyQM6DggAEIAEELEDEIMBEJIDUMcJWJYyYOQ9aAJwAHgAgAGsAYgBzAqSAQQxMi4ymAEAoAEBsAEK&sclient=gws-wiz'
        convert = Currency.currency_soup(DOLLAR_BUN).findAll("span", {"class": "DFlfde SwHCTb", "data-precision": 2})
        result = convert[0].text
        return result

    @classmethod
    def currency_rub(cls):
        RUB_BUN = 'https://www.google.by/search?q=курс+рубля&sxsrf=APwXEdfg26YNygnwvIva9NmX64eKvpkdQQ%3A1684148019264&source=hp&ei=Mw9iZMzvDIaF9u8P5PSPsAE&iflsig=AOEireoAAAAAZGIdQ4ZxDKL8tZppf6yFKlHHnxgzK_Di&oq=курс+&gs_lcp=Cgdnd3Mtd2l6EAEYADIHCCMQigUQJzIHCCMQigUQJzIHCCMQigUQJzILCAAQgAQQsQMQgwEyCwgAEIAEELEDEIMBMgsIABCABBCxAxCDATILCAAQgAQQsQMQgwEyBQgAEIAEMgsIABCABBCxAxCDATIFCAAQgAQ6BwgjEOoCECc6CwgAEIoFELEDEIMBOg4ILhCABBCxAxCDARDUAjoFCC4QgAQ6CAguEIAEENQCOggILhDUAhCABDoKCAAQgAQQyQMQCjoICAAQigUQkgM6CAgAEIAEELEDOg4IABCABBCxAxCDARDJAzoOCAAQgAQQsQMQgwEQkgNQb1iZDGDBFWgBcAB4AYABwgGIAZwFkgEDNi4xmAEAoAEBsAEK&sclient=gws-wiz'
        convert = Currency.currency_soup(RUB_BUN).findAll("span", {"class": "DFlfde SwHCTb", "data-precision": 3})
        result = convert[0].text
        return result

class Search_program():
    paths = open(r"C:\Users\mrgod\PycharmProjects\VoiceAsistent\Paths.txt", "a+")
    ready_path = ''

    @classmethod
    def search_program(cls, program):
        Assistant.talk("Сейчас попробую найти на вашем компьютере\n Это может занять время\n но в следующий раз будет быстрее")
        for root, dirs, files in os.walk('C:\\'):
            #print ("searching", root)
            if program in files:
                path = "%s" % join(root, program)
                Search_program.paths.write(f"\n{path}")
                break

    @classmethod
    def search_in_txt(cls, program) -> str:
        with io.open(r"C:\Users\mrgod\PycharmProjects\VoiceAsistent\Paths.txt", encoding='utf-8') as file:
            lines = [line.rstrip() for line in file]
            for line in lines:
                if program in line:
                    print("файл найден")
                    ready_path = line
                    return ready_path
            Search_program.search_program(program)

if __name__ == "__main__":
    app = QtWidgets.QApplication(sys.argv)
    window = Assistant()
    window.show()
    sys.exit(app.exec_())